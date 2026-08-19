import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_krishisetu_presentation_with_images():
    prs = Presentation()
    prs.slide_width = Inches(13.333) # 16:9 Widescreen
    prs.slide_height = Inches(7.5)

    BASE_DIR = os.path.dirname(os.path.abspath(__file__))
    LOGO_PATH = os.path.join(BASE_DIR, "src", "assets", "logos", "logo.jpeg")
    
    # Check logo image existence
    has_logo = os.path.exists(LOGO_PATH)

    # Color Palette
    PRIMARY_GREEN = RGBColor(27, 94, 32)    # #1B5E20
    SECONDARY_GREEN = RGBColor(67, 160, 71)  # #43A047
    ACCENT_GOLD = RGBColor(249, 168, 37)    # #F9A825
    DARK_TEXT = RGBColor(31, 41, 55)       # #1F2937
    LIGHT_BG = RGBColor(248, 250, 245)      # #F8FAF5
    WHITE = RGBColor(255, 255, 255)

    blank_slide_layout = prs.slide_layouts[6]

    slides_data = [
        {
            "title": "KRISHISETU",
            "subtitle": "Bridging Farmers & Consumers\n\nA Production-Grade Agriculture Marketplace Platform",
            "bullets": [
                "Tagline: Bridging Farmers & Consumers",
                "Frontend: React 19 + Vite + Bootstrap 5 + Framer Motion",
                "Backend: Django REST Framework + SimpleJWT + MySQL",
                "Presenter: Senior Full-Stack Engineer"
            ],
            "is_hero": True,
            "show_logo": True
        },
        {
            "title": "1. Problem Statement & Market Need",
            "subtitle": "Challenges in Traditional Agricultural Supply Chains",
            "bullets": [
                "Middleman Exploitation: Traditional commission agents absorb up to 60% of crop value.",
                "Low Farmer Returns: Kisans receive unfair prices despite intense physical labor.",
                "Lack of Harvest Transparency: Urban consumers pay inflated prices for non-traceable produce.",
                "Delayed Payments: Farmers suffer from delayed payment cycles and complex market logistics."
            ]
        },
        {
            "title": "2. Proposed Solution: KRISHISETU",
            "subtitle": "Direct Farm-to-Consumer Digital Marketplace Ecosystem",
            "bullets": [
                "Zero Middleman Commissions: Farmers set their own crop prices and keep up to 96.5% of revenue.",
                "100% Traceable Harvests: Every crop listing includes harvest timestamp and farm location.",
                "Direct Farm Dispatches: Fresh produce harvested morning 5:30 AM delivered within 24h.",
                "Multi-Portal Governance: Integrated ecosystems for Buyers, Farmers, and Super Admin officers."
            ],
            "show_logo": True
        },
        {
            "title": "3. System Architecture & Tech Stack",
            "subtitle": "Modern Commercial E-Commerce Stack",
            "bullets": [
                "Frontend Shell: React 19, Vite, React Router DOM v6, Axios, React Toastify.",
                "UI/UX Design Tokens: Custom Green/Gold Design System, Framer Motion, Recharts analytics.",
                "Backend Infrastructure: Django 5.x REST Framework, SimpleJWT Auth, django-cors-headers.",
                "Database Tier: MySQL Relational Database (`krishisetu_db`) with PyMySQL / mysqlclient drivers."
            ]
        },
        {
            "title": "4. Database Schema & Entity Relationships",
            "subtitle": "Relational Data Modeling in MySQL",
            "bullets": [
                "accounts_user: Custom AbstractUser with roles ('buyer', 'farmer', 'admin').",
                "accounts_farmerprofile: Farm estate name, district, state, specialties JSON, rating, verification status.",
                "productsapp_product: Name, category, price, unit ('kg', 'dozen'), stock, harvest_date, is_organic, specs.",
                "ordersapp_order & tracking: Order number, payment status, delivery address, live progress timeline steps."
            ]
        },
        {
            "title": "5. Public Marketplace Features",
            "subtitle": "High-Impact Commercial E-Commerce Experience",
            "bullets": [
                "Full-Width 64px Hero Section: 600x500px agricultural banner with green announcement badge.",
                "1-Click Category Chips: Quick horizontal harvest filters (🥦 Veggies, 🍎 Fruits, 🌾 Grains, 🥛 Dairy).",
                "Slide-Over Quick Cart Drawer: Instant cart preview, quantity modification, and coupon validator.",
                "Full-Screen Logo Watermark: Official KrishiSetu emblem watermark fitted to 100% viewport background."
            ],
            "show_logo": True
        },
        {
            "title": "6. Farmer Producer Portal Features",
            "subtitle": "Empowering Rural Kisan Producers",
            "bullets": [
                "1-Click Crop Preset Templates: Preset high-res harvest images (Tomatoes, Alphonso Mangoes, Wheat, Ghee, Honey).",
                "Crop Catalog Management: Real-time price updates, available stock tracking, and organic badges.",
                "Order Dispatch Controls: Farmers view incoming orders and trigger harvest dispatch status.",
                "Recharts Sales Analytics: Monthly revenue growth charts, quintals sold, and payout summaries."
            ],
            "show_logo": True
        },
        {
            "title": "7. Super Admin Governance Portal",
            "subtitle": "Platform Integrity & Verification Control",
            "bullets": [
                "Kisan Verification Engine: Review and approve newly registered farmer profiles.",
                "Marketplace Quality Audit: Inspect crop listings, prices, and organic compliance.",
                "Platform Metrics: Real-time system transaction totals, active buyer counts, and revenue statistics.",
                "User Governance: Account suspension and policy enforcement tools."
            ]
        },
        {
            "title": "8. Security & System Integration",
            "subtitle": "Enterprise-Grade API Integration",
            "bullets": [
                "JWT Token Authentication: 7-day access token lifetime with auto Axios Bearer interceptor.",
                "Role-Based Authorization Guards: DRF custom permissions (`IsFarmerOrAdminOrReadOnly`).",
                "Stock Deduction Logic: Automatic stock reduction upon successful order placement.",
                "CORS Security: Strict origin permission configured for React Vite (`http://localhost:5173`)."
            ]
        },
        {
            "title": "9. Real-Time Order Tracking Timeline",
            "subtitle": "5-Step Farm Dispatch Progress Tracking",
            "bullets": [
                "Step 1: Order Placed (Instant Order Number Generation).",
                "Step 2: Harvested & Packed (Farm estate harvest confirmation).",
                "Step 3: Dispatched from Farm (Regional agri-hub dispatch).",
                "Step 4: Out for Delivery (Urban doorstep delivery vehicle).",
                "Step 5: Delivered (Customer receipt confirmation)."
            ]
        },
        {
            "title": "10. Future Enhancements & Roadmap",
            "subtitle": "Scaling KrishiSetu Nationwide",
            "bullets": [
                "IoT Farm Sensors: Integrating soil moisture & climate sensors directly into Kisan profiles.",
                "AI Yield Prediction: Machine learning crop price forecasting and seasonal yield advice.",
                "Voice-Based Multilingual Search: Voice search supporting Hindi, Marathi, Punjabi, and Kannada.",
                "Cold-Chain Fleet Logistics: Real-time GPS temperature tracking for refrigerated produce trucks."
            ]
        },
        {
            "title": "Thank You! Q & A",
            "subtitle": "KRISHISETU - Bridging Farmers & Consumers",
            "bullets": [
                "Project Name: KRISHISETU",
                "Repository: krishisetu-frontend & backend",
                "Technology: React 19 + Vite + Django REST + MySQL",
                "Questions & Discussion Welcome!"
            ],
            "is_hero": True,
            "show_logo": True
        }
    ]

    for data in slides_data:
        slide = prs.slides.add_slide(blank_slide_layout)

        # Background shape
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        shape.fill.solid()
        if data.get("is_hero"):
            shape.fill.fore_color.rgb = PRIMARY_GREEN
        else:
            shape.fill.fore_color.rgb = LIGHT_BG
        shape.line.fill.background()

        # Regular Slide Header
        if not data.get("is_hero"):
            top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.2))
            top_bar.fill.solid()
            top_bar.fill.fore_color.rgb = PRIMARY_GREEN
            top_bar.line.fill.background()

            txBox = slide.shapes.add_textbox(Inches(0.6), Inches(0.15), Inches(12), Inches(0.9))
            tf = txBox.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = data["title"]
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = WHITE

            txSub = slide.shapes.add_textbox(Inches(0.6), Inches(1.3), Inches(12), Inches(0.5))
            tfSub = txSub.text_frame
            pSub = tfSub.paragraphs[0]
            pSub.text = data["subtitle"]
            pSub.font.size = Pt(18)
            pSub.font.bold = True
            pSub.font.color.rgb = SECONDARY_GREEN
        else:
            # Hero Slide Title
            txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(7.5), Inches(1.5))
            tf = txBox.text_frame
            p = tf.paragraphs[0]
            p.text = data["title"]
            p.font.size = Pt(50)
            p.font.bold = True
            p.font.color.rgb = ACCENT_GOLD

            txSub = slide.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(7.5), Inches(1.2))
            tfSub = txSub.text_frame
            pSub = tfSub.paragraphs[0]
            pSub.text = data["subtitle"]
            pSub.font.size = Pt(20)
            pSub.font.color.rgb = WHITE

        # Image Embedding on Slides
        if data.get("show_logo") and has_logo:
            if data.get("is_hero"):
                # Embedded Logo Image Card on Hero / Title Slide
                slide.shapes.add_picture(LOGO_PATH, Inches(8.5), Inches(1.5), width=Inches(4.2))
            else:
                # Embedded Logo Image Card on Feature Slides
                slide.shapes.add_picture(LOGO_PATH, Inches(9.2), Inches(2.0), width=Inches(3.5))

        # Bullets Body Container
        if not data.get("is_hero"):
            width = Inches(8.2) if (data.get("show_logo") and has_logo) else Inches(11.8)
            body_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), width, Inches(5.0))
            tf_body = body_box.text_frame
            tf_body.word_wrap = True

            for bullet in data["bullets"]:
                p = tf_body.add_paragraph()
                p.text = "•  " + bullet
                p.font.size = Pt(19)
                p.font.color.rgb = DARK_TEXT
                p.space_after = Pt(14)
        else:
            body_box = slide.shapes.add_textbox(Inches(0.8), Inches(4.2), Inches(7.5), Inches(2.8))
            tf_body = body_box.text_frame
            tf_body.word_wrap = True

            for bullet in data["bullets"]:
                p = tf_body.add_paragraph()
                p.text = "✓  " + bullet
                p.font.size = Pt(18)
                p.font.color.rgb = WHITE
                p.space_after = Pt(10)

    output_path = os.path.join(BASE_DIR, "KRISHISETU_Final_Presentation.pptx")
    prs.save(output_path)
    print(f"Final presentation with embedded images generated successfully at: {output_path}")

if __name__ == "__main__":
    build_krishisetu_presentation_with_images()
